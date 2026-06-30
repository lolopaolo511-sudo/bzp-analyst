"""Testy ulepszeń 4-6: synonimy, SearchText pre-filtr, eksport Excel."""
import sys
import io
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent.parent))

import pytest
from datetime import date, timedelta
from unittest.mock import patch, MagicMock

from sources.normalizer import normalize_notice
from workflow.scorer import (
    SYNONYMS,
    _expand_with_synonyms,
    _keyword_score,
    _kw_in_text,
)
from workflow.pipeline import WorkflowConfig, run_pipeline


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _raw_notice(title: str, cpv: str = "79952000", cpv_desc: str = "Organizacja imprez",
                days_left: int = 14) -> dict:
    deadline = (date.today() + timedelta(days=days_left)).strftime("%Y-%m-%dT12:00:00Z")
    return {
        "objectId": f"test-{abs(hash(title)) % 100000}",
        "noticeNumber": f"2026/BZP/{abs(hash(title)) % 10000}/01",
        "noticeType": "ContractNotice",
        "orderType": "Services",
        "orderObject": title,
        "organizationName": "Urząd Gminy Test",
        "organizationCity": "Kraków",
        "organizationProvince": "PL21",
        "organizationCountry": "PL",
        "cpvCode": f"{cpv} ({cpv_desc})",
        "publicationDate": "2026-06-01T08:00:00Z",
        "submittingOffersDate": deadline,
        "isTenderAmountBelowEU": True,
    }


# ---------------------------------------------------------------------------
# [4] Synonimy — słownik i rozszerzanie
# ---------------------------------------------------------------------------

class TestSynonymsDict:
    def test_synonyms_is_dict(self):
        assert isinstance(SYNONYMS, dict)

    def test_key_naglośnienie_has_synonyms(self):
        assert "nagłośnienie" in SYNONYMS
        assert len(SYNONYMS["nagłośnienie"]) >= 2

    def test_key_impreza_has_synonyms(self):
        assert "impreza" in SYNONYMS
        assert "event" in SYNONYMS["impreza"]

    def test_key_koncert_has_synonyms(self):
        assert "koncert" in SYNONYMS

    def test_key_scena_has_synonyms(self):
        assert "scena" in SYNONYMS

    def test_all_synonym_values_are_lists(self):
        for k, v in SYNONYMS.items():
            assert isinstance(v, list), f"Synonimy dla '{k}' nie są listą"

    def test_all_synonyms_are_lowercase(self):
        for k, synonyms in SYNONYMS.items():
            for s in synonyms:
                assert s == s.lower(), f"Synonim '{s}' dla '{k}' nie jest lowercase"


class TestExpandWithSynonyms:
    def test_expands_known_keyword(self):
        result = _expand_with_synonyms(["impreza"])
        assert "impreza" in result
        assert "event" in result

    def test_unknown_keyword_unchanged(self):
        result = _expand_with_synonyms(["xyzabc"])
        assert result == ["xyzabc"]

    def test_no_duplicates_in_expansion(self):
        result = _expand_with_synonyms(["impreza", "impreza"])
        assert len(result) == len(set(result)), "Rozszerzenie zawiera duplikaty"

    def test_multiple_keywords_expanded(self):
        result = _expand_with_synonyms(["nagłośnienie", "koncert"])
        assert "nagłośnienie" in result
        assert "koncert" in result
        # Synonimy z obu
        assert any(s in result for s in SYNONYMS.get("nagłośnienie", []))
        assert any(s in result for s in SYNONYMS.get("koncert", []))

    def test_empty_list_returns_empty(self):
        assert _expand_with_synonyms([]) == []


class TestSynonymScoring:
    EVENTS_KW = ["nagłośnienie", "impreza", "koncert", "scena"]

    def test_synonym_event_matches_impreza_keyword(self):
        """'event' w tytule powinien dać score > 0 gdy keyword 'impreza'."""
        text = "Obsługa event plenerowy w parku miejskim"
        score = _keyword_score(text, ["impreza"], [])
        assert score > 0, "Synonim 'event' nie trafił dla keyword 'impreza'"

    def test_synonym_estrada_matches_scena_keyword(self):
        text = "Wynajem estrada na festyn"
        score = _keyword_score(text, ["scena"], [])
        assert score > 0

    def test_synonym_recital_matches_koncert(self):
        text = "Organizacja recital muzycznego w filharmonii"
        score = _keyword_score(text, ["koncert"], [])
        assert score > 0

    def test_synonym_akustyka_matches_naglośnienie(self):
        text = "Projekt akustyki sali widowiskowej"
        score = _keyword_score(text, ["nagłośnienie"], [])
        assert score > 0

    def test_synonym_does_not_break_exclude(self):
        """Synonim słowa wykluczonego nie może przepuścić ogłoszenia."""
        text = "Obsługa event sportowy"
        # "impreza" i jego synonimy — ale wykluczamy "sportowy"
        score = _keyword_score(text, ["impreza"], ["sportowy"])
        assert score == 0.0

    def test_synonym_score_at_most_1(self):
        text = "Nagłośnienie impreza koncert scena event festyn estrada"
        score = _keyword_score(text, self.EVENTS_KW, [])
        assert 0.0 <= score <= 1.0

    def test_direct_keyword_still_matches(self):
        """Normalne słowa kluczowe nadal działają (regresja)."""
        text = "Nagłośnienie imprezy plenerowej w parku"
        score = _keyword_score(text, ["nagłośnienie", "impreza"], [])
        assert score == 1.0

    def test_no_match_synonym_or_direct_gives_zero(self):
        text = "Roboty budowlane drogi krajowej nr 7"
        score = _keyword_score(text, ["nagłośnienie", "impreza"], [])
        assert score == 0.0


# ---------------------------------------------------------------------------
# [5] SearchText pre-filtr — pipeline konfiguracja
# ---------------------------------------------------------------------------

class TestSearchTextPreFilter:
    def test_pipeline_config_has_use_api_search_text(self):
        cfg = WorkflowConfig()
        assert hasattr(cfg, "use_api_search_text")

    def test_use_api_search_text_default_false(self):
        # Domyślnie False — BZP SearchText przeszukuje pełny tekst dokumentu
        # (nie tylko tytuł), co daje masę false-positives. Filtrujemy lokalnie.
        cfg = WorkflowConfig()
        assert cfg.use_api_search_text is False

    def test_pipeline_passes_search_text_when_enabled(self):
        """Gdy use_api_search_text=True, pipeline przekazuje keyword jako search_text do BZPQuery."""
        captured_queries = []

        def mock_fetch_for_keyword(client, kw, *args, **kwargs):
            captured_queries.append(kw)
            return []

        with patch("workflow.pipeline._fetch_for_keyword", side_effect=mock_fetch_for_keyword):
            with patch("workflow.pipeline.BZPClient"):
                run_pipeline(WorkflowConfig(
                    keywords_include=["nagłośnienie", "impreza"],
                    use_api_search_text=True,
                    use_llm=False,
                ))

        assert "nagłośnienie" in captured_queries
        assert "impreza" in captured_queries

    def test_pipeline_uses_empty_keyword_when_disabled(self):
        """Gdy use_api_search_text=False, pipeline pobiera bez filtra tekstowego."""
        captured_queries = []

        def mock_fetch_for_keyword(client, kw, *args, **kwargs):
            captured_queries.append(kw)
            return []

        with patch("workflow.pipeline._fetch_for_keyword", side_effect=mock_fetch_for_keyword):
            with patch("workflow.pipeline.BZPClient"):
                run_pipeline(WorkflowConfig(
                    keywords_include=["nagłośnienie", "impreza"],
                    use_api_search_text=False,
                    use_llm=False,
                ))

        assert captured_queries == [""], "Powinno być jedno puste zapytanie"

    def test_pipeline_uses_empty_when_no_keywords(self):
        """Brak słów kluczowych → jedno puste zapytanie bez SearchText."""
        captured_queries = []

        def mock_fetch_for_keyword(client, kw, *args, **kwargs):
            captured_queries.append(kw)
            return []

        with patch("workflow.pipeline._fetch_for_keyword", side_effect=mock_fetch_for_keyword):
            with patch("workflow.pipeline.BZPClient"):
                run_pipeline(WorkflowConfig(
                    keywords_include=[],
                    use_api_search_text=True,
                    use_llm=False,
                ))

        assert captured_queries == [""]

    def test_fetch_for_keyword_passes_search_text_to_query(self):
        """_fetch_for_keyword przekazuje search_text do BZPQuery."""
        from workflow.pipeline import _fetch_for_keyword
        mock_client = MagicMock()
        mock_client.fetch_all.return_value = iter([])

        _fetch_for_keyword(mock_client, "nagłośnienie", ["ContractNotice"], 7, [], 5)

        call_args = mock_client.fetch_all.call_args
        query = call_args[0][0]
        assert query.search_text == "nagłośnienie"

    def test_fetch_for_keyword_empty_string_gives_none_search_text(self):
        """Puste słowo kluczowe → search_text=None (bez filtra w API)."""
        from workflow.pipeline import _fetch_for_keyword
        mock_client = MagicMock()
        mock_client.fetch_all.return_value = iter([])

        _fetch_for_keyword(mock_client, "", ["ContractNotice"], 7, [], 5)

        call_args = mock_client.fetch_all.call_args
        query = call_args[0][0]
        assert query.search_text is None


# ---------------------------------------------------------------------------
# [6] Eksport Excel — generowanie pliku XLSX
# ---------------------------------------------------------------------------

class TestExcelExport:
    def _make_df(self):
        import pandas as pd
        return pd.DataFrame([
            {
                "Nr BZP": "2026/BZP 00001/01",
                "Score (0-1)": 0.75,
                "Składowe": "CPV 0.90 · KW 0.70 · DL 1.00 · VAL 0.50",
                "Tytuł": "Nagłośnienie imprezy plenerowej",
                "Zamawiający": "Urząd Gminy Kraków",
                "Miasto": "Kraków",
                "Województwo": "Małopolskie",
                "Termin (dni)": 14,
                "Typ": "Ogłoszenie o zamówieniu",
                "CPV": "79952000-2",
                "Wartość (PLN)": 150000.0,
                "Link BZP": "https://ezamowienia.gov.pl/mo-client-board/bzp/notice-details/test-001",
            }
        ])

    def _to_excel(self, df):
        import io
        import pandas as pd
        buf = io.BytesIO()
        with pd.ExcelWriter(buf, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Przetargi BZP")
            ws = writer.sheets["Przetargi BZP"]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
        return buf.getvalue()

    def test_excel_output_is_bytes(self):
        df = self._make_df()
        result = self._to_excel(df)
        assert isinstance(result, bytes)

    def test_excel_output_nonempty(self):
        df = self._make_df()
        result = self._to_excel(df)
        assert len(result) > 1000  # plik XLSX ma nagłówek + dane

    def test_excel_starts_with_xlsx_magic(self):
        df = self._make_df()
        result = self._to_excel(df)
        # XLSX to ZIP — zaczyna się od PK\x03\x04
        assert result[:2] == b"PK"

    def test_excel_readable_by_openpyxl(self):
        import io
        import openpyxl
        df = self._make_df()
        xlsx_bytes = self._to_excel(df)
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        assert "Przetargi BZP" in wb.sheetnames

    def test_excel_sheet_has_correct_columns(self):
        import io
        import openpyxl
        df = self._make_df()
        xlsx_bytes = self._to_excel(df)
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        ws = wb["Przetargi BZP"]
        headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
        assert "Tytuł" in headers
        assert "Score (0-1)" in headers
        assert "Link BZP" in headers

    def test_excel_has_correct_row_count(self):
        import io
        import openpyxl
        import pandas as pd
        df = pd.concat([self._make_df(), self._make_df()], ignore_index=True)
        xlsx_bytes = self._to_excel(df)
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        ws = wb["Przetargi BZP"]
        # +1 dla nagłówka
        assert ws.max_row == len(df) + 1

    def test_excel_freeze_panes_set(self):
        import io
        import openpyxl
        df = self._make_df()
        xlsx_bytes = self._to_excel(df)
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        ws = wb["Przetargi BZP"]
        assert ws.freeze_panes == "A2"

    def test_excel_empty_dataframe(self):
        """Pusty DataFrame — plik XLSX nadal ważny."""
        import pandas as pd
        empty_df = pd.DataFrame(columns=["Tytuł", "Score (0-1)"])
        result = self._to_excel(empty_df)
        assert isinstance(result, bytes) and len(result) > 0
